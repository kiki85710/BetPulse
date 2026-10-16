"""Construit le classeur planning_betpulse.xlsx à partir de epreuves.csv.

Usage : python planning/build_planning.py
Le fichier .xlsx produit est ensuite importé dans Google Sheets.
"""
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ICI = Path(__file__).parent
SOURCE = ICI / "epreuves.csv"
SORTIE = ICI / "planning_betpulse.xlsx"

# Référentiel des compétitions : code, sport, compétition, granularité, source officielle
COMPETITIONS = [
    ("L1", "Football", "Ligue 1", "1 ligne par journée", "https://ligue1.com"),
    ("L2", "Football", "Ligue 2", "1 ligne par journée", "https://www.ligue2.fr"),
    ("LIGA", "Football", "LaLiga", "1 ligne par journée", "https://www.laliga.com"),
    ("EPL", "Football", "Premier League", "1 ligne par journée", "https://www.premierleague.com"),
    ("UCL", "Football", "UEFA Champions League", "1 ligne par journée ou par match aller/retour", "https://www.uefa.com/uefachampionsleague"),
    ("UNL", "Football", "UEFA Nations League", "1 ligne par journée ou par tour", "https://www.uefa.com/uefanationsleague"),
    ("T14", "Rugby", "Top 14", "1 ligne par journée", "https://top14.lnr.fr"),
    ("PD2", "Rugby", "Pro D2", "1 ligne par journée", "https://prod2.lnr.fr"),
    ("CC", "Rugby", "Champions Cup", "1 ligne par journée ou par tour", "https://www.epcrugby.com"),
    ("SL", "Handball", "StarLigue", "1 ligne par journée", "https://www.lnh.fr"),
    ("MAG", "Hockey sur glace", "Ligue Magnus", "1 ligne par journée ou par tour de play-offs", "https://www.hockeyfrance.com"),
    ("SPK", "Volley-ball", "Alterna Volstar", "1 ligne par journée ou par tour de play-offs", "https://www.lnv.fr"),
    ("NBA", "Basketball", "NBA", "1 ligne par semaine + temps forts", "https://www.nba.com"),
    ("NFL", "Football américain", "NFL", "1 ligne par semaine (Week)", "https://www.nfl.com"),
    ("PGA", "Golf", "PGA Tour", "1 ligne par tournoi", "https://www.pgatour.com"),
    ("F1", "Sport automobile", "Formule 1", "1 ligne par week-end de Grand Prix", "https://www.formula1.com"),
    ("MGP", "Moto", "MotoGP", "1 ligne par week-end de Grand Prix", "https://www.motogp.com"),
    ("WRC", "Rallye", "WRC", "1 ligne par rallye", "https://www.wrc.com"),
    ("GC", "Tennis", "Grand Chelem", "1 ligne par tournoi", "https://www.itftennis.com"),
    ("UCI", "Cyclisme sur route", "UCI WorldTour", "1 ligne par course", "https://www.uci.org"),
    ("UCIM", "Cyclisme sur route", "Championnats du monde UCI", "1 ligne par épreuve", "https://www.uci.org"),
    ("SKI", "Ski alpin", "Coupe du monde de ski alpin", "1 ligne par épreuve", "https://www.fis-ski.com"),
    ("SKIM", "Ski alpin", "Championnats du monde de ski alpin", "1 ligne par épreuve", "https://www.fis-ski.com"),
    ("BIA", "Biathlon", "Coupe du monde de biathlon", "1 ligne par épreuve", "https://www.biathlonworld.com"),
    ("BIAM", "Biathlon", "Championnats du monde de biathlon", "1 ligne par épreuve", "https://www.biathlonworld.com"),
    ("EHFCL", "Handball", "EHF Champions League", "1 ligne par journée ou par tour", "https://www.eurohandball.com"),
    ("EHFEL", "Handball", "EHF European League", "1 ligne par journée ou par tour", "https://www.eurohandball.com"),
    ("CDFH", "Handball", "Coupe de France de handball", "1 ligne par tour", "https://www.ffhandball.fr"),
    ("TDCH", "Handball", "Trophée des Champions de handball", "1 ligne par édition", "https://www.lnh.fr"),
    ("IHFSG", "Handball", "IHF Super Globe", "1 ligne par édition", "https://www.ihf.info"),
    ("MCH", "Handball", "Championnat du monde de handball", "1 ligne par édition", "https://www.ihf.info"),
    ("QEH", "Handball", "Qualifications Euro de handball", "1 ligne par fenêtre internationale", "https://www.eurohandball.com"),
]
CODE = {c[2]: c[0] for c in COMPETITIONS}
SPORTS = sorted({c[1] for c in COMPETITIONS})
STATUTS = ["Officiel", "Provisoire", "À confirmer", "Reporté", "Annulé"]
GENRES = ["Hommes", "Femmes", "Mixte"]
OUI_NON = ["Oui", "Non"]
DATE_VERIF = date(2026, 10, 16)

COLONNES = [
    ("id", 16, "Identifiant unique et stable : CODE-SAISON-NUMÉRO. Ne jamais le modifier une fois créé (le code Python s'en sert comme clé)."),
    ("sport", 18, "Sport, choisi dans la liste (onglet Référentiels)."),
    ("competition", 30, "Compétition, choisie dans la liste (onglet Référentiels)."),
    ("saison", 10, "Saison : 2026-2027 pour les sports d'hiver et les championnats à cheval sur deux années, 2026 ou 2027 pour les calendriers sur une année civile."),
    ("phase", 22, "Phase de la compétition : saison régulière, phase de ligue, play-offs, finale, discipline (ski, biathlon)..."),
    ("manche", 9, "Numéro de journée, de semaine ou de manche. Vide si la phase n'est pas numérotée."),
    ("intitule", 44, "Libellé lisible de l'épreuve, tel qu'il sera affiché dans BetPulse."),
    ("date_debut", 12, "Premier jour de l'épreuve (format JJ/MM/AAAA)."),
    ("date_fin", 12, "Dernier jour de l'épreuve. Identique à date_debut pour une épreuve d'un jour."),
    ("heure_debut", 11, "Heure de départ, heure de Paris (HH:MM). Renseignée seulement si elle est officielle."),
    ("fuseau", 13, "Fuseau horaire de heure_debut. Toujours Europe/Paris."),
    ("lieu", 32, "Site et ville. Vide pour une journée de championnat (plusieurs stades)."),
    ("pays", 14, "Pays hôte."),
    ("genre", 9, "Hommes, Femmes ou Mixte."),
    ("statut", 13, "Officiel : date confirmée par l'organisateur. Provisoire : publiée mais susceptible de bouger (programmation TV, calendrier provisoire). À confirmer : pas encore publiée. Reporté / Annulé."),
    ("ouvert_pronos", 13, "Oui si l'épreuve est proposée aux pronostics dans BetPulse, sinon Non."),
    ("source_url", 40, "Page web sur laquelle la date a été vérifiée."),
    ("date_verification", 13, "Date de la dernière vérification de la ligne."),
    ("notes", 40, "Précision utile : week-end sprint, match délocalisé, programme détaillé non publié..."),
    ("avancement", 12, "Calculé automatiquement (ne pas saisir) : À venir, En cours ou Terminé selon la date du jour."),
]

POLICE = "Arial"
BLEU = "1F3864"
GRIS = "D9D9D9"
fin = Side(style="thin", color="BFBFBF")
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)
ENTETE_FONT = Font(name=POLICE, bold=True, color="FFFFFF", size=10)
ENTETE_FILL = PatternFill("solid", fgColor=BLEU)
CORPS = Font(name=POLICE, size=10)


def lire_epreuves():
    with SOURCE.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def entete(ws, ligne, valeurs, largeurs=None):
    for j, v in enumerate(valeurs, start=1):
        c = ws.cell(row=ligne, column=j, value=v)
        c.font, c.fill, c.border = ENTETE_FONT, ENTETE_FILL, BORD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if largeurs:
        for j, w in enumerate(largeurs, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def onglet_lisez_moi(wb, nb_lignes):
    ws = wb.active
    ws.title = "Lisez-moi"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 110
    titre = ws.cell(row=2, column=2, value="Planning des épreuves BetPulse")
    titre.font = Font(name=POLICE, bold=True, size=18, color=BLEU)
    ws.cell(row=3, column=2, value=f"Saison 2026-2027 et suivantes. Données vérifiées sur le web le {DATE_VERIF:%d/%m/%Y}.").font = Font(name=POLICE, italic=True, size=10, color="595959")

    blocs = [
        ("À quoi sert ce fichier", [
            "Il recense toutes les épreuves proposées aux pronostics dans BetPulse : une ligne par épreuve, toutes compétitions confondues.",
            "L'application (ou un script Python) lit l'onglet Épreuves pour savoir quelles épreuves ouvrir aux pronostics et à quelles dates.",
        ]),
        ("Les onglets", [
            "Épreuves : la table de données. C'est le seul onglet à mettre à jour au quotidien.",
            "Référentiels : la liste des sports, compétitions, statuts et genres autorisés, avec le nombre d'épreuves et la prochaine date par compétition.",
            "Lisez-moi : cette notice.",
        ]),
        ("Granularité retenue", [
            "Sports mécaniques (F1, MotoGP, WRC) : une ligne par week-end de course ou par rallye.",
            "Sports collectifs (football, rugby, handball, hockey, volley, NFL) : une ligne par journée. NBA : une ligne par semaine, plus les temps forts (NBA Cup, All-Star, play-offs, Finales).",
            "Ski alpin et biathlon : une ligne par épreuve (une course = une ligne, hommes et femmes séparés).",
            "Golf, tennis, cyclisme : une ligne par tournoi ou par course (les courses par étapes tiennent sur une ligne, du premier au dernier jour).",
        ]),
        ("Règles de saisie", [
            "1. Une ligne = une épreuve. Ne pas fusionner de cellules, ne pas laisser de ligne vide au milieu de la table.",
            "2. Les colonnes sport, competition, statut, genre et ouvert_pronos ont une liste déroulante : choisir une valeur de la liste.",
            "3. Les dates se saisissent au format JJ/MM/AAAA. Les heures sont en heure de Paris (HH:MM).",
            "4. L'identifiant (colonne id) ne se modifie jamais. Pour une nouvelle épreuve : CODE-SAISON-NUMÉRO, le code étant celui de l'onglet Référentiels (ex. F1-2027-025).",
            "5. Une épreuve annulée ou reportée n'est pas supprimée : changer son statut et préciser la raison dans notes.",
            "6. À chaque vérification, mettre à jour source_url et date_verification.",
            "7. La colonne avancement se calcule seule. Pour une nouvelle ligne, recopier la formule de la ligne du dessus (dans Google Sheets, elle s'étend toute seule).",
        ]),
        ("Code couleur", [
            "Ligne grisée : épreuve terminée.",
            "Statut en orange : date provisoire. Statut en rouge : date à confirmer, reportée ou annulée.",
        ]),
        ("Lecture en Python", [
            "Voir planning/lire_planning.py dans le dépôt BetPulse. Le script lit l'onglet Épreuves (export CSV ou API Google Sheets) et renvoie un DataFrame pandas.",
        ]),
    ]
    r = 5
    for titre_bloc, lignes in blocs:
        c = ws.cell(row=r, column=2, value=titre_bloc)
        c.font = Font(name=POLICE, bold=True, size=12, color=BLEU)
        r += 1
        for l in lignes:
            c = ws.cell(row=r, column=3, value=l)
            c.font = CORPS
            c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1

    c = ws.cell(row=r, column=2, value="Dictionnaire des colonnes (onglet Épreuves)")
    c.font = Font(name=POLICE, bold=True, size=12, color=BLEU)
    r += 1
    for j, v in enumerate(["Colonne", "Contenu"], start=2):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.fill, c.border = ENTETE_FONT, ENTETE_FILL, BORD
    r += 1
    for nom, _, desc in COLONNES:
        a = ws.cell(row=r, column=2, value=nom)
        a.font = Font(name=POLICE, size=10, bold=True)
        b = ws.cell(row=r, column=3, value=desc)
        b.font = CORPS
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = b.border = BORD
        r += 1


def onglet_epreuves(wb, rows):
    ws = wb.create_sheet("Épreuves")
    entete(ws, 1, [c[0] for c in COLONNES], [c[1] for c in COLONNES])
    for i, e in enumerate(rows, start=2):
        valeurs = [
            e["id"], e["sport"], e["competition"], e["saison"], e["phase"],
            int(e["manche"]) if e["manche"].isdigit() else (e["manche"] or None),
            e["intitule"],
            date.fromisoformat(e["date_debut"]) if e["date_debut"] else None,
            date.fromisoformat(e["date_fin"]) if e["date_fin"] else None,
            e["heure_debut"] or None, "Europe/Paris", e["lieu"] or None, e["pays"] or None,
            e["genre"], e["statut"], e.get("ouvert_pronos") or "Oui",
            e["source_url"] or None, DATE_VERIF, e["notes"] or None,
            f'=IF(H{i}="","",IF(MAX(H{i},I{i})<TODAY(),"Terminé",IF(H{i}<=TODAY(),"En cours","À venir")))',
        ]
        for j, v in enumerate(valeurs, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = CORPS
            c.border = BORD
            c.alignment = Alignment(vertical="center")
        for col in (8, 9, 18):
            ws.cell(row=i, column=col).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=10).number_format = "@"
    der = len(rows) + 1
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{der}"
    ws.row_dimensions[1].height = 30

    # Listes déroulantes alimentées par l'onglet Référentiels, sur une plage qui dépasse la table
    fin_dv = der + 500
    for col, plage in [("B", "'Référentiels'!$H$2:$H$30"), ("C", "'Référentiels'!$C$2:$C$60"),
                       ("O", "'Référentiels'!$J$2:$J$10"), ("N", "'Référentiels'!$K$2:$K$10"),
                       ("P", "'Référentiels'!$L$2:$L$5")]:
        dv = DataValidation(type="list", formula1=f"={plage}", allow_blank=True, showErrorMessage=True,
                            errorTitle="Valeur non autorisée", error="Choisir une valeur de la liste (onglet Référentiels).")
        dv.add(f"{col}2:{col}{fin_dv}")
        ws.add_data_validation(dv)
    for col in ("H", "I", "R"):
        dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                            showErrorMessage=True, errorTitle="Date attendue", error="Saisir une date au format JJ/MM/AAAA.")
        dv.add(f"{col}2:{col}{fin_dv}")
        ws.add_data_validation(dv)

    plage = f"A2:T{fin_dv}"
    ws.conditional_formatting.add(plage, FormulaRule(formula=['$T2="Terminé"'], font=Font(color="808080"), fill=PatternFill("solid", fgColor="F2F2F2")))
    ws.conditional_formatting.add(f"O2:O{fin_dv}", FormulaRule(formula=['$O2="Provisoire"'], font=Font(color="9C5700", bold=True), fill=PatternFill("solid", fgColor="FFEB9C")))
    ws.conditional_formatting.add(f"O2:O{fin_dv}", FormulaRule(formula=['OR($O2="À confirmer",$O2="Reporté",$O2="Annulé")'], font=Font(color="9C0006", bold=True), fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(f"O2:O{fin_dv}", FormulaRule(formula=['$O2="Officiel"'], font=Font(color="006100"), fill=PatternFill("solid", fgColor="C6EFCE")))
    return der


def onglet_referentiels(wb, der):
    ws = wb.create_sheet("Référentiels")
    entete(ws, 1, ["code", "sport", "competition", "granularite", "source_officielle", "nb_epreuves", "prochaine_epreuve", "sports", "", "statuts", "genres", "oui_non"],
           [8, 18, 34, 40, 38, 12, 16, 20, 3, 13, 10, 9])
    for i, (code, sport, comp, gran, src) in enumerate(COMPETITIONS, start=2):
        for j, v in enumerate([code, sport, comp, gran, src], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = CORPS, BORD
        f = ws.cell(row=i, column=6, value=f"=COUNTIF('Épreuves'!$C$2:$C$5000,C{i})")
        g = ws.cell(row=i, column=7, value=f"=IFERROR(1/(1/_xlfn.MINIFS('Épreuves'!$H$2:$H$5000,'Épreuves'!$C$2:$C$5000,C{i},'Épreuves'!$H$2:$H$5000,\">=\"&TODAY())),\"\")")
        g.number_format = "DD/MM/YYYY"
        for c in (f, g):
            c.font, c.border = CORPS, BORD
    tot = len(COMPETITIONS) + 2
    t = ws.cell(row=tot, column=3, value="Total")
    t.font = Font(name=POLICE, bold=True, size=10)
    s = ws.cell(row=tot, column=6, value=f"=SUM(F2:F{tot - 1})")
    s.font = Font(name=POLICE, bold=True, size=10)
    for col, liste in [(8, SPORTS), (10, STATUTS), (11, GENRES), (12, OUI_NON)]:
        for i, v in enumerate(liste, start=2):
            c = ws.cell(row=i, column=col, value=v)
            c.font, c.border = CORPS, BORD
    ws.cell(row=1, column=9).fill = PatternFill(fill_type=None)
    ws.freeze_panes = "A2"


def main():
    rows = lire_epreuves()
    wb = Workbook()
    onglet_lisez_moi(wb, len(rows))
    der = onglet_epreuves(wb, rows)
    onglet_referentiels(wb, der)
    wb.save(SORTIE)
    print(f"{len(rows)} épreuves -> {SORTIE}")


if __name__ == "__main__":
    main()
